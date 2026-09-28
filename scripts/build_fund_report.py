"""Build the Portuguese NORTE report from frozen evidence. No provider requests.

Requires system matplotlib, python-docx (installed separately outside the repo),
and LibreOffice for the subsequent PDF conversion. Does not change benchmark data.
"""
import sys
sys.path.insert(0, '/tmp/norte-report-libs')
import json
import math
import textwrap
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/gama-fund-2026-09-27'
FIG = OUT / 'figuras'
FIG.mkdir(parents=True, exist_ok=True)
A = json.loads((OUT / 'metricas_agregadas.json').read_text())
AN = json.loads((ROOT / '.runtime/fund-report/analysis.json').read_text())
NAVY, BLUE, GREEN, AMBER, RED = '#15324B', '#3575D3', '#16836B', '#AB740D', '#C44848'
GRAY, LIGHT = '#526779', '#F0F5F8'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'text.color':NAVY,'axes.labelcolor':NAVY,
                     'xtick.color':GRAY,'ytick.color':NAVY,'axes.edgecolor':'#D7E1E8','figure.facecolor':'white'})

def group(model, protocol, task, level='L1_DIRECT'):
    return next(x for x in A if x['model']==model and x['comparison_hash'].startswith(protocol) and x['task']==task and x['difficulty']==level)

def metric(g, name):
    return g['metrics'][name]['mean']

def pct(x):
    return '—' if x is None else f'{100*x:.1f}%'.replace('.',',')

def num(x, digits=1):
    return f'{x:.{digits}f}'.replace('.',',')

def save(fig, name):
    fig.savefig(FIG/f'{name}.png',dpi=210,bbox_inches='tight',facecolor='white')
    fig.savefig(FIG/f'{name}.svg',bbox_inches='tight',facecolor='white')
    plt.close(fig)

def box(ax, x,y,w,h,label,color=BLUE,fill=LIGHT,fontsize=10):
    p=FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.015,rounding_size=0.04',linewidth=1.3,edgecolor=color,facecolor=fill)
    ax.add_patch(p);ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=fontsize,color=NAVY)

def arrow(ax,start,end,color=BLUE,dashed=False):
    ax.add_patch(FancyArrowPatch(start,end,arrowstyle='-|>',mutation_scale=12,color=color,linewidth=1.8,linestyle='--' if dashed else '-'))

fig,ax=plt.subplots(figsize=(10,3.35));ax.set(xlim=(0,10),ylim=(0,3.5));ax.axis('off')
for x,title in [(0.05,'5 PDFs originais\n120 páginas'),(2.62,'Extração local\ntexto por página'),(5.19,'API do modelo\numa tarefa por chamada'),(7.76,'Resposta estruturada\nJSON + evidências')]:
    box(ax,x,2.12,2.15,1,title)
for x in [2.22,4.79,7.36]:arrow(ax,(x,2.62),(x+.35,2.62))
box(ax,2.62,.28,2.15,.95,'Requisitos e escolhas\ndo projeto',GRAY)
arrow(ax,(4.8,.77),(5.8,2.1),GRAY)
box(ax,5.19,.28,2.15,.95,'Gabarito provisório\nfora do prompt',AMBER)
box(ax,7.76,.28,2.15,.95,'Validação + comparação\nrevisão humana',GREEN)
arrow(ax,(8.83,2.1),(8.83,1.25),GREEN);arrow(ax,(7.38,.77),(7.73,.77),GREEN)
ax.text(.12,.58,'Tarefas independentes.\nNão é uma cadeia de saídas\nalimentando outros modelos.',fontsize=9,color=GRAY)
save(fig,'01_metodologia')

series=[('Gemini 3.1 Pro\nb43c00 · n=3','gemini-3.1-pro-preview','b43c00',BLUE),
        ('GPT-6 Astra\n678f30 · n=3','gpt-6-astra','678f30',GREEN),
        ('Claude Fable\nb43c00 · n=1','claude-fable-5-1','b43c00',AMBER)]
fig,axes=plt.subplots(1,3,figsize=(10,3.3),sharey=True)
for ax,(key,title) in zip(axes,[('relationship_precision','Precisão'),('relationship_recall','Cobertura (recall)'),('relationship_f1','F1 de relações')]):
    for i,(label,model,h,c) in enumerate(series):
        g=group(model,h,'relationship_extraction');v=metric(g,key)*100
        ax.barh(i,v,color=c,height=.44);ax.text(v+1.5,i,f'{v:.1f}%'.replace('.',','),va='center',fontsize=10)
    ax.set(xlim=(0,118),xticks=[0,50,100],yticks=range(3),yticklabels=[s[0] for s in series],title=title)
    ax.grid(axis='x',alpha=.15);ax.set_axisbelow(True);ax.spines[['top','right','left']].set_visible(False)
axes[0].invert_yaxis();fig.tight_layout();save(fig,'02_relacoes')

fig,axes=plt.subplots(1,2,figsize=(10,4.55))
for ax,title in zip(axes,['Esperado · sistema instalado','Recebido · Gemini, repetição 1']):
    ax.set(xlim=(0,5),ylim=(0,5));ax.axis('off');ax.set_title(title,loc='left',fontsize=12,fontweight='bold')
    box(ax,.12,3.65,1.32,.62,'REQ-001\nLimite de corrente',GRAY,fontsize=9)
    box(ax,.12,2.05,1.32,.62,'REQ-002\nVelocidade mínima',GRAY,fontsize=9)
    box(ax,.12,.48,1.32,.62,'FAN instalado\nNF-A4x10',GRAY,fontsize=9)
ax=axes[0]
box(ax,3.16,3.65,1.63,.62,'P-FAN-CURRENT\nCorrente instalada',GREEN,fontsize=9)
box(ax,3.16,2.05,1.63,.62,'P-FAN-SPEED\nVelocidade instalada',GREEN,fontsize=9)
box(ax,3.16,.48,1.63,.62,'V-SPEED\nTeste do projeto',GREEN,fontsize=9)
for y,label in [(3.96,'limita'),(2.36,'limita')]:arrow(ax,(1.5,y),(3.1,y),GREEN);ax.text(2.3,y+.14,label,ha='center',fontsize=8,color=GREEN)
arrow(ax,(1.46,.9),(3.13,3.75),GREEN);arrow(ax,(1.46,.8),(3.13,2.13),GREEN)
ax.text(1.8,1.08,'possui parâmetros',fontsize=8,color=GREEN,rotation=52)
arrow(ax,(1.48,2.15),(3.13,.82),GREEN);ax.text(2.4,1.15,'verificado por',fontsize=8,color=GREEN,rotation=-35)
ax=axes[1]
box(ax,3.16,3.65,1.63,.62,'P-ALT-FAN-CURRENT\nCorrente alternativa',RED,fontsize=8.5)
box(ax,3.16,2.05,1.63,.62,'P-ALT-FAN-SPEED\nVelocidade alternativa',RED,fontsize=8.5)
box(ax,3.16,.48,1.63,.62,'V-SPEED\nTeste do projeto',GREEN,fontsize=9)
for y in [3.96,2.36]:arrow(ax,(1.5,y),(3.1,y),RED);ax.text(2.3,y+.14,'alvo divergente',ha='center',fontsize=8,color=RED)
ax.text(1.68,2.93,'As ligações ao ventilador\ninstalado foram omitidas.',fontsize=9,color=RED)
arrow(ax,(1.48,2.15),(3.13,.82),GREEN);ax.text(2.4,1.15,'verificado por',fontsize=8,color=GREEN,rotation=-35)
ax.text(.13,.04,'Vermelho: erro neste recorte do sistema base.\nVerde: relação coincide com a referência.',fontsize=8.5,color=GRAY)
fig.tight_layout();save(fig,'03_diagrama_gemini')

tasks=[('change_impact','L1_DIRECT','Impactos L1'),('impact_explanation','L1_DIRECT','Explicações L1'),('impact_explanation','L2_ONE_HOP','Explicações L2'),('one_hop','L2_ONE_HOP','Inferência L2')]
fig,ax=plt.subplots(figsize=(10,3.3))
for i,(task,level,label) in enumerate(tasks):
    g=group('gemini-3.1-pro-preview','b43c00',task,level)
    for delta,key,c in [(-.16,'impact_f1',GREEN),(.16,'explanation_rule_pass_rate',AMBER)]:
        v=metric(g,key)*100;ax.barh(i+delta,v,height=.26,color=c);ax.text(v+1,i+delta,f'{v:.1f}%'.replace('.',','),va='center',fontsize=10)
ax.set(yticks=range(4),yticklabels=[t[2] for t in tasks],xlim=(0,113),xticks=[0,25,50,75,100]);ax.invert_yaxis();ax.grid(axis='x',alpha=.16);ax.set_axisbelow(True)
ax.spines[['top','right','left']].set_visible(False)
ax.plot([],[],lw=7,color=GREEN,label='F1: quais requisitos foram afetados');ax.plot([],[],lw=7,color=AMBER,label='Justificativas que passam todas as regras')
ax.legend(loc='lower center',bbox_to_anchor=(.5,1.04),ncol=1,frameon=False);fig.tight_layout();save(fig,'04_impactos_e_justificativas')

alltasks=[('entity_extraction','L1_DIRECT','Extração'),('relationship_extraction','L1_DIRECT','Relações')]+tasks
fig,ax=plt.subplots(figsize=(10,3.8))
for i,(task,level,label) in enumerate(alltasks):
    for offset,model,h,c in [(-.17,'gemini-3.1-pro-preview','b43c00',BLUE),(.17,'gpt-6-astra','678f30',GREEN)]:
        g=group(model,h,task,level);s=g['latency_seconds'];ax.barh(i+offset,s['mean'],height=.28,color=c,alpha=.9)
        ax.errorbar(s['mean'],i+offset,xerr=[[s['mean']-s['min']],[s['max']-s['mean']]],color=NAVY,fmt='none',capsize=2,lw=.9)
        ax.text(s['max']+2,i+offset,num(s['mean'])+' s',va='center',fontsize=9)
ax.set(yticks=range(6),yticklabels=[t[2] for t in alltasks],xlim=(0,205),xlabel='Segundos por resposta avaliada · barras: média; linhas: mínimo–máximo')
ax.invert_yaxis();ax.spines[['top','right','left']].set_visible(False);ax.grid(axis='x',alpha=.15);ax.set_axisbelow(True)
ax.plot([],[],lw=7,color=BLUE,label='Gemini · b43c00');ax.plot([],[],lw=7,color=GREEN,label='Astra · 678f30')
ax.legend(loc='lower right',frameon=False);fig.tight_layout();save(fig,'05_latencia')

# The editable document is the source of the PDF, so captions and numbers match.
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.55);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
styles=doc.styles
for name in ['Normal','Body Text']:
    styles[name].font.name='Calibri';styles[name].font.size=Pt(10.5)
    styles[name].font.color.rgb=RGBColor.from_string('243D50')
    styles[name].paragraph_format.space_after=Pt(7);styles[name].paragraph_format.line_spacing=1.1
for name,size in [('Title',30),('Heading 1',21),('Heading 2',13),('Heading 3',11)]:
    styles[name].font.name='Calibri';styles[name].font.size=Pt(size);styles[name].font.bold=True
    styles[name].font.color.rgb=RGBColor.from_string('15324B')
styles['Caption'].font.size=Pt(8.5);styles['Caption'].font.color.rgb=RGBColor.from_string('526779')
header=sec.header.paragraphs[0];header.text='NORTE  /  EVIDÊNCIAS DE DESENVOLVIMENTO';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('27.09.2026  •  Estudo exploratório  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.style='Caption'
doc.core_properties.title='NORTE — Onde os modelos Google ajudam e onde ficam aquém'
doc.core_properties.subject='Evidências para a questão 5 do formulário Gama Fund'
doc.core_properties.author='Projeto NORTE'
md=[]
next_page=False

def title(t):
    doc.add_heading(t,0);md.append('# '+t+'\n')
def h(t):
    global next_page
    paragraph=doc.add_heading(t,1)
    if next_page:paragraph.paragraph_format.page_break_before=True
    next_page=False
    md.append('## '+t+'\n')
def sub(t):
    doc.add_heading(t,2);md.append('### '+t+'\n')
def p(t):
    doc.add_paragraph(t);md.append(t+'\n')
def bullet(t):
    doc.add_paragraph(t,style='List Bullet');md.append('- '+t+'\n')
def note(t):
    x=doc.add_paragraph(t);x.paragraph_format.space_before=Pt(5)
    for r in x.runs:r.font.color.rgb=RGBColor.from_string('805C0A')
    md.append('> '+t+'\n')
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
    for c,txt in zip(t.rows[0].cells,headers):c.text=str(txt)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(4);pp.paragraph_format.line_spacing=1
                for rr in pp.runs:rr.font.size=Pt(9)
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    if widths:
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Cm(w)
    md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |'])
    md.extend('| '+' | '.join(map(str,row))+' |' for row in rows);md.append('')
def pic(name,caption,width=17.3):
    doc.add_picture(str(FIG/f'{name}.png'),width=Cm(width));doc.add_paragraph(caption,style='Caption')
    md.append(f'![{caption}](figuras/{name}.png)\n')
def page():
    global next_page
    next_page=True
    md.append('\n---\n')

title('Modelos de IA para\nrastreabilidade de engenharia')
sub('NORTE · Relatório de evidências para a questão 5')
p('Onde os modelos do Google se saem bem e onde ficam aquém no processo de desenvolvimento?')
p('Preparado para acompanhar a resposta ao formulário do Gama Fund. Análise de testes locais realizados em 27 de setembro de 2026, com corte dos registros às 22h33 (America/São_Paulo). Relatório e gráficos derivados de respostas reais já salvas; nenhuma nova chamada paga foi feita para elaborar este documento.')
sub('O principal achado')
p('O Gemini 3.1 Pro Preview foi útil para transformar documentação técnica extensa em informação estruturada e identificar quais requisitos precisam ser revistos após uma mudança. Completou as 18 avaliações previstas e recuperou todos os requisitos afetados nos cenários testados. Ainda não oferece, neste protótipo, a precisão e a rastreabilidade necessárias para aceitar automaticamente o grafo e suas justificativas.')
table(['Evidência do Gemini','Resultado observado'],[
('Execução completa','18 de 18 respostas avaliadas; todas parcialmente corretas pelas regras salvas'),
('Identificação de impactos','F1 de 100% nas quatro combinações de tarefa/nível de impacto; n=3 em cada'),
('Relações entre itens','Precisão 55,1%; cobertura 81,8%; F1 65,0%'),
('Justificativas completas','Aprovação automática de 0% a 8,3%, conforme a tarefa'),
('Latência','66,8 s em média nas 18 respostas; relações: 79,9 s')])
sub('Como usar a conclusão')
p('O resultado sustenta o uso do Gemini como assistente de análise documental e triagem, com inspeção humana das relações e evidências. Não sustenta certificação de engenharia, autonomia irrestrita nem superioridade geral de um fornecedor. Apenas um modelo Google foi testado neste recorte.')
note('Leitura essencial: o gabarito é provisório, há restrições do avaliador que rejeitam citações reais e os protocolos não são todos idênticos. Os percentuais são sinais de desenvolvimento, não um ranking definitivo. O Astra d4408d, declarado inválido, foi excluído.')

page();h('1. O caso de uso e o teste conduzido')
p('O NORTE explora rastreabilidade entre componentes, parâmetros, requisitos e verificações. O objetivo é permitir que uma pessoa examine o que foi extraído, quais relações foram propostas, que requisitos uma alteração afeta e em quais documentos essas conclusões se apoiam. Este recorte avalia modelos integrados ao produto; não mede a qualidade deles como assistentes de programação.')
pic('01_metodologia','Figura 1. Fluxo efetivamente testado. O PDF é a fonte original; a API recebeu texto extraído, instruções e um esquema de resposta. O gabarito foi reservado ao avaliador.')
sub('Material de entrada')
p('Foram usados cinco PDFs de fabricantes: Noctua NF-A4x10 5V (7 páginas), TI TMP117 (50), TI TPS22919 (27), Noctua NF-A4x20 5V (7) e TI TPS22917 (29). São 120 páginas de origem. Os dois últimos documentos fornecem evidência para substituições propostas. A extração local produziu texto identificado por página; não houve ensaio de leitura visual nativa dos PDFs pelos modelos.')
p('A esse material foram acrescentados cinco requisitos e cinco trechos de definição do projeto. Alimentações de 5 V e 3,3 V, limiar de 35 °C e controle liga/desliga são escolhas declaradas do projeto. Não são dados extraídos pelo modelo nem medições físicas. O sistema conceitual usa ventilador, sensor, chave de alimentação e um controlador ESP32; nenhum hardware foi construído ou qualificado neste teste.')
sub('Escopo da referência')
p('A referência salva contém 17 entidades, 13 itens numéricos, 11 relações e 12 cenários de mudança. Há nove cenários L1, incluindo dois controles sem impacto, e três cenários L2 de inferência curta. Os rótulos foram criados durante a implementação, com assistência de IA, antes das respostas dos provedores. A revisão técnica humana independente continua pendente: 0 de 58 itens ratificados no snapshot.')

page();h('2. Procedimento e controle de versões')
table(['Tarefa','Nível','Pergunta avaliada'],[
('Extração de entidades e números','L1','Quais itens, valores, unidades e fontes foram identificados?'),
('Extração de relações','L1','Quem possui um parâmetro, limita um valor, depende de uma configuração ou é verificado por um teste?'),
('Identificação de impactos','L1','Quais requisitos precisam de revisão em cada mudança?'),
('Explicação dos impactos','L1','A dependência e a justificativa do impacto têm suporte?'),
('Explicação dos impactos','L2','A justificativa se mantém nas três mudanças com inferência curta?'),
('Inferência de um salto','L2','O modelo deduz os impactos a partir dos documentos e da mudança?')])
p('O plano previa três repetições de cada combinação, totalizando 18 respostas por modelo. Uma resposta de impacto reúne vários cenários; 18 respostas não significam 18 projetos independentes. Cada tarefa recebeu seu próprio prompt e os documentos, sem encadear a saída de extração como entrada da tarefa de relações. Todos os resultados deste relatório são de primeira passagem, sem feedback confirmado no contexto.')
table(['Modelo registrado na API','Configuração salva','Amostra elegível'],[
('gemini-3.1-pro-preview','thinking_level=high; temperatura 1','18 respostas · b43c00'),
('gpt-6-astra','reasoning_effort=high; temperatura não fixada','17 respostas · 678f30'),
('claude-fable-5-1','reasoning_effort=high; temperatura não fixada','8 respostas · três versões')])
p('Configuração: limite de 16.000 tokens de saída, timeout de 300 segundos e retries=1. Os nomes acima são os registrados nas chamadas. Os níveis “high” não equivalem a um orçamento de computação controlado entre APIs.')
sub('O que mudou entre versões')
p('b43c00 e 678f30 preservam o mesmo dataset, os mesmos prompts 2.0.0 e os mesmos esquemas, mas têm hashes diferentes do código avaliador. 7f007f acrescenta o contrato de unicidade dos impactos no prompt 2.0.1 e possui esquema/código diferentes. As médias permanecem separadas por protocolo. A comparação entre modelos é descritiva, sem reavaliar retrospectivamente as respostas.')
note('Exclusão: as 18 respostas Astra d4408d (protocolo histórico 1.1.0 / prompt 1.0.0) foram invalidadas pela responsável pelo estudo. Permanecem no histórico; seus antigos 100% não são usados neste relatório.')

page();h('3. Google: pontos fortes na extração')
p('O Gemini recuperou as 17 entidades e os 13 identificadores numéricos esperados em todas as três repetições. Entre os itens numéricos que coincidiram com a referência, acertou os valores em 92,3% e as unidades em 100%. A assinatura de extração foi igual nas três repetições, indicando estabilidade desse conjunto de decisões, embora uma resposta estável também possa repetir um erro.')
table(['Indicador · extração L1, n=3','Média','O que significa'],[
('Cobertura de entidades / parâmetros','100% / 100%','Nenhum identificador esperado foi omitido'),
('F1 de entidades / parâmetros','82,9% / 83,9%','Itens extras reduzem a precisão'),
('Exatidão dos valores / unidades','92,3% / 100%','Somente fatos com ID reconhecido'),
('Atribuição de fonte','61,9%','Combina localização de entidades e evidência numérica'),
('Cobertura de fatos integralmente aceitos','76,9%','Exige valor, unidade, qualificador e evidência aceitos')])
sub('Onde a extração ficou aquém')
p('Cada resposta trouxe 24 entidades para uma referência de 17, e 18 fatos numéricos para uma referência de 13. Parte dos extras descreve componentes alternativos presentes nos documentos. Isso é inadequado para o grafo do sistema instalado, mas não equivale automaticamente a inventar um fato. A separação entre sistema base e alternativas de mudança precisa ficar mais clara no contexto, no esquema e na avaliação.')
table(['Exemplo real','Recebido do Gemini','Esperado / interpretação'],[
('P-FAN-VOLTAGE','Faixa de 4,0 a 5,5 V; qualificador range','O ID da referência designa tensão nominal: 5 V, rated. São atributos distintos do mesmo PDF.'),
('P-SENSOR-SUPPLY','1,8 a 5,5 V, citando SENSOR página 1','Valor correto e trecho literal. A regra aceita apenas a página 5 para esse atributo.'),
('P-DRIVER-VOLTAGE','1,6 a 5,5 V, citando DRIVER página 1','Valor correto e trecho literal. A regra aceita apenas a página 4 para esse atributo.')])
note('Os dois últimos casos são uma limitação demonstrável do avaliador: os trechos existem no texto preservado da página 1. As notas históricas foram mantidas, mas não devem ser descritas como alucinações do Gemini. A revisão futura deve aceitar fontes equivalentes com suporte técnico.')
p('Para o desenvolvimento, a utilidade está em reduzir a busca inicial e estruturar uma primeira versão inspecionável. Não foi medido quanto tempo de trabalho humano essa ajuda economiza; esse benefício ainda precisa de validação com usuários.')

page();h('4. Relações: cobertura útil, precisão insuficiente')
pic('02_relacoes','Figura 2. Médias por resposta na extração de relações. Gemini e Astra: três respostas cada; Claude: uma. Protocolos identificados; não há teste de significância nem ranking controlado.')
p('O Gemini teve F1 médio de relações de 65,0%, precisão de 55,1% e cobertura de 81,8%. Identificou uma parcela relevante das dependências, mas aproximadamente metade das relações propostas não coincidia com a referência ou não passava o critério de evidência. A taxa média chamada “unsupported_relationship_rate” foi 53,3%; ela reúne relações extras e relações corretas com citação reprovada, inclusive o caso de página equivalente descrito anteriormente.')
table(['Repetição Gemini','Corretas / extras / ausentes','F1','Corretas com evidência aceita'],[
('1','5 / 6 / 6','45,5%','4 de 11 propostas'),
('2','11 / 6 / 0','78,6%','10 de 17 propostas'),
('3','11 / 9 / 0','71,0%','9 de 20 propostas')])
p('O F1 que também exige evidência aprovada cai para 55,3% em média. Nas três respostas, os conjuntos completos de arestas foram diferentes; a consistência por igualdade exata foi 0%. Isso não significa que nenhuma relação se repetiu: significa que nenhum par de respostas reproduziu o grafo inteiro.')
sub('O que os outros modelos acrescentam à leitura')
p('O Astra 678f30 recuperou todas as 11 relações nas três respostas e apoiou essas relações com evidências aceitas. Porém, acrescentou mais arestas: precisão média de 43,1% e F1 de 60,2%. O Claude b43c00 teve precisão de 52,4%, cobertura de 100% e F1 de 68,8%, mas em uma única resposta. Essa amostra não permite concluir que Claude é superior.')
note('A taxa de itens extras depende de um gabarito de escopo fechado. Uma relação plausível, porém fora desse escopo, pode receber a mesma marca de uma relação realmente incorreta. Antes de publicar um ranking, é necessário revisar a completude do grafo e classificar as divergências.')

page();h('5. Um erro visível no grafo do Gemini')
p('Na primeira repetição de relações, o Gemini associou os limites de corrente e velocidade aos parâmetros do ventilador alternativo NF-A4x20, em vez do ventilador instalado NF-A4x10. Os documentos alternativos estavam presentes para analisar mudanças futuras. O prompt de relações pedia o sistema base.')
pic('03_diagrama_gemini','Figura 3. Recorte fiel de relações da resposta 397e93c4… (Gemini, repetição 1). O desenho simplifica a disposição dos nós; os IDs e as ligações vêm da resposta e da referência salvas. Não representa o grafo completo.')
table(['Relação','Referência do sistema base','Resposta nessa repetição'],[
('Limite de corrente','REQ-001 → constrains → P-FAN-CURRENT','REQ-001 → constrains → P-ALT-FAN-CURRENT'),
('Velocidade mínima','REQ-002 → constrains → P-FAN-SPEED','REQ-002 → constrains → P-ALT-FAN-SPEED'),
('Verificação de velocidade','REQ-002 → verified_by → V-SPEED','Coincide com a referência e tem evidência aceita')])
p('Esse caso mostra uma falha relevante para o produto: reconhecer fatos de um documento não basta; é necessário associá-los à versão correta do sistema. Se uma alternativa de componente entrar silenciosamente no grafo instalado, a análise de mudanças perde confiabilidade.')
p('Nas repetições 2 e 3, as relações do ventilador instalado reapareceram, mas as ligações extras às alternativas permaneceram. A melhoria de cobertura entre repetições, portanto, não resolveu a necessidade de distinguir contexto atual e contexto proposto.')
sub('Consequência para a interface e o fluxo')
p('Uma pessoa precisa conseguir ver o alvo recebido e o alvo esperado, filtrar as relações por tipo e decidir se há erro do modelo, ambiguidade de escopo ou lacuna na referência. A evidência justifica revisão por item, com status visível; uma porcentagem geral não explica esse problema.')

page();h('6. Identificar impactos não é justificar impactos')
pic('04_impactos_e_justificativas','Figura 4. Gemini, três respostas por tarefa. O verde mede identificação de requisitos afetados; o amarelo mede justificativas que passam simultaneamente todas as regras estruturadas. São medidas diferentes.')
p('O Gemini acertou os conjuntos de requisitos afetados em todos os cenários e repetições das quatro tarefas de impacto. Não houve falsos positivos de impacto, omissões críticas ou impacto indevido nos dois cenários de controle sem mudança relevante. O resultado é promissor para triagem documental e localização do que precisa ser revisto.')
p('A aprovação das justificativas foi muito menor: 0 de 24 na tarefa de impactos L1; 2 de 24 nas explicações L1; 1 de 12 nas explicações L2; 0 de 12 na inferência L2. Esses denominadores contam impactos previstos nas respostas, repetidos entre tarefas; não são casos independentes. A regra exige requisito, elemento alterado, dependência, citações e afirmações numéricas corretos, além de uma explicação não vazia.')
sub('Exemplo: substituição do ventilador, CHG-001')
table(['Campo · REQ-001, explicações L1, repetição 1','Resultado'],[
('Requisito afetado','Recebido REQ-001; coincide com a referência'),
('Componente e dependência','FAN; REQ-001 → constrains → P-FAN-CURRENT; corretos'),
('Afirmação numérica','0,05 A, rated, com citação do PDF; aceita'),
('Evidência da dependência','O campo source_evidence cita só a mudança CHG-001; falta a fonte base exigida nesse campo'),
('Avaliação salva','Impacto identificado corretamente; justificativa reprovada pela regra de evidência')])
p('A fonte do fato numérico aparece dentro de technical_claims, mas não substitui a citação da dependência no campo principal. Isso é um problema de atendimento ao contrato de evidências e também uma decisão rígida do avaliador. Não autoriza concluir que a explicação inteira é fisicamente falsa.')
note('F1 de impacto de 100% não significa 100% de respostas corretas. As 18 respostas Gemini foram classificadas como parciais. A qualidade da prosa não recebeu avaliação humana independente, e nenhuma nota certifica o funcionamento de hardware.')

page();h('7. Latência e experiência de desenvolvimento')
pic('05_latencia','Figura 5. Latência das respostas avaliadas. Gemini: n=3 em cada tarefa. Astra: n=3, exceto explicações L1, n=2. As séries usam versões diferentes do código avaliador e não foram executadas simultaneamente.')
p('A média global do Gemini foi 66,8 segundos por resposta, variando de 35,4 a 97,4 segundos. Na tarefa de relações, foram 79,9 segundos, contra 155,1 segundos do Astra 678f30. Esse é um sinal favorável de tempo de resposta neste ambiente, não uma garantia universal de velocidade. Os tamanhos das saídas, tokenizadores, ajustes de raciocínio e condições do serviço diferem.')
sub('O que isso significa para a experiência')
p('Para uma avaliação executada em segundo plano, essa latência pode ser administrada com progresso visível, resultados parciais e retomada. Para um fluxo de edição que exige resposta imediata, esperar dezenas de segundos a cada chamada ainda interrompe a interação. O estudo não definiu um SLA nem mediu tolerância de usuários; por isso não classifica 79 segundos como “bom” em termos absolutos.')
table(['Uso registrado · Gemini, 18 respostas','Total'],[
('Tokens de entrada','1.181.154'),('Tokens de saída','182.972'),('Custos monetários','Não disponíveis nos registros; nenhum valor foi estimado')])
p('O contexto de cada chamada Gemini tinha aproximadamente 65–66 mil tokens de entrada. O mesmo conjunto documental foi usado em várias tarefas; as contagens incluem reutilização do conteúdo, e há registros de cache em algumas respostas. Não se deve converter esses totais diretamente em custo nem comparar tokenizadores como se fossem a mesma unidade de texto.')
sub('Melhoria sugerida para o produto')
p('Vale testar recuperação de trechos relevantes, cache e reutilização de extrações verificadas, mantendo acesso à fonte original e medindo o efeito sobre a cobertura. Também vale comparar tamanhos de saída menores e níveis de raciocínio. São experimentos propostos: nenhum ganho de tempo ou custo dessas mudanças foi demonstrado neste estudo.')

page();h('8. Claude: problema operacional e de contrato')
p('A experiência com Claude foi irregular, mas os registros não mostram ausência total de respostas. Há oito respostas avaliadas distribuídas por três versões e quatro chamadas com erro técnico explícito. O histórico também contém tarefas interrompidas, que não devem ser tratadas como respostas erradas do modelo nem como requisições necessariamente enviadas.')
table(['Falha observada','Evidência','Interpretação'],[
('Saldo/faturamento','1 chamada com HTTP 400 e billing_error','A API bloqueou essa chamada. O saldo exibido em outra tela não permite comprovar a origem da divergência de conta/chave naquele momento.'),
('Contrato da resposta','3 chamadas com HTTP 200, encerramento end_turn e requisito duplicado no cenário','A API respondeu, mas o benchmark rejeitou a estrutura semântica; não foi uma medição de qualidade concluída.'),
('Tentativa mais recente','1 resposta one_hop L2 em 7f007f, avaliada após prompt 2.0.1','F1 de impacto 100%; aprovação de justificativas 0%. Uma resposta não demonstra resolução geral.')])
sub('Exemplo exato da rejeição')
table(['Cenário CHG-004','Requisito','Dependência recebida'],[
('Entrada 1','REQ-002','depends_on → P-FAN-SPEED'),
('Entrada 2','REQ-002','verified_by → V-SPEED')])
p('O contrato aceitava apenas uma entrada por requisito em cada cenário. A resposta apresentou duas dependências em dois registros do mesmo requisito. A validação parou com duplicate_impact; não foi um estouro comprovado do limite de saída. As outras duas falhas repetem REQ-002 em CHG-012.')
p('O prompt antigo não explicitava suficientemente essa unicidade. A versão 2.0.1 passou a orientar uma entrada por requisito e uma dependência principal, preservando as relações adicionais na explicação. A tentativa posterior foi aceita estruturalmente, mas não aprovou suas justificativas. Portanto, houve progresso de integração sem evidência de confiabilidade completa.')
sub('Onde Claude mostrou capacidade')
p('Nas respostas avaliadas de impacto, identificou os requisitos esperados. Na única resposta válida de relações b43c00, recuperou as 11 relações da referência com evidências aceitas. Ainda assim, a amostra é pequena e seletiva: as falhas interromperam parte do plano. Não é adequado atribuir nota zero de raciocínio a billing, nem usar só a resposta bem-sucedida para declarar superioridade.')

page();h('9. Síntese por modelo e limites de interpretação')
table(['Modelo','Onde se destacou neste recorte','Onde ficou aquém'],[
('Gemini 3.1 Pro Preview','Completou as 18 avaliações; cobertura integral de entidades e impactos; menor latência observada que Astra por tarefa.','Relações extras/instáveis, confusão entre base e alternativas, diferença entre valor nominal e faixa, evidências incompletas ou rejeitadas.'),
('GPT-6 Astra · 678f30','Cobertura de 100% das relações; relações esperadas com evidência aceita; cobertura numérica integral, 92,3% de fatos integralmente aceitos na extração.','Excesso de relações reduz precisão para 43,1%; explicações com baixa aprovação; latências maiores neste registro; amostra de 17 respostas.'),
('Claude Fable · versões separadas','Nas respostas válidas, localizou impactos; única resposta de relações recuperou todas as arestas esperadas.','Uma falha billing e três falhas de contrato; oito respostas avaliadas em três versões; insuficiente para uma comparação equilibrada.')])
sub('O que estes testes não medem')
bullet('Produtividade humana, qualidade de código gerado, satisfação de usuários ou economia financeira. Não há medição antes/depois nem custo monetário completo.')
bullet('Leitura nativa de imagens, tabelas visuais e diagramas de PDF. Os modelos receberam texto extraído, que pode perder a organização visual das páginas.')
bullet('Generalização para projetos grandes, outros domínios, documentos contraditórios ou centenas de componentes. Trata-se de um único sistema conceitual e do mesmo conjunto de cenários repetido.')
bullet('Aprendizado com feedback. Todos os registros usados são first_pass e não têm correções confirmadas no contexto. Os cenários marcados transfer não demonstram aprendizado adquirido entre chamadas.')
bullet('Validade física de uma solução ou compatibilidade elétrica/mecânica de uma substituição. A referência pede revisão da evidência de aceitação, não certifica aprovação ou reprovação do hardware.')
note('Além do gabarito não ratificado, há um possível desalinhamento de escopo: a referência é um grafo fechado, enquanto os prompts atuais pedem descobrir fatos aplicáveis sem fornecer todos os IDs esperados. Fontes equivalentes e afirmações corretas fora do escopo devem ser revisadas antes de chamar toda divergência de erro factual.')
p('As comparações são exploratórias. Não há intervalo de confiança confiável com três repetições de um único caso, nem ensaio cego independente. As notas foram lidas do histórico sem substituir saídas por gabarito ou corrigir automaticamente respostas.')

page();h('10. O que melhorar a partir das evidências')
table(['Prioridade','Melhoria proposta','Como verificar'],[
('1 · Avaliação','Ratificar o gabarito com pessoa tecnicamente competente; aceitar citações equivalentes; separar fora de escopo, tipo errado, fonte ausente e fato falso.','Revisar os 58 itens e os exemplos página 1/5/4; versionar a regra e manter resultados antigos.'),
('2 · Contexto do modelo','Separar sistema instalado de alternativas; explicitar atributos como tensão nominal versus faixa operacional; exigir fonte para cada relação.','Casos sem dicas da resposta, com alternativas distratoras; medir precisão sem perder cobertura.'),
('3 · Contrato de saída','Validar unicidade, IDs e campos de evidência antes da avaliação; instruções e esquema devem expressar a mesma regra.','Executar casos de duplicação e relações múltiplas; reportar rejeição estrutural separada de erro semântico.'),
('4 · Desempenho','Reduzir contexto repetido e tamanho de saída; experimentar recuperação de fontes e cache.','Medir latência mediana/p95, tokens, custo real e qualidade no mesmo protocolo.'),
('5 · Confiabilidade','Comparar os modelos em uma versão congelada e executar apenas tentativas explicitamente selecionadas.','Mesmas entradas/regras, plano completo, falhas contabilizadas, histórico preservado.'),
('6 · Uso por pessoas','Mostrar saída × esperado, diferenças nas arestas, fonte original e feedback por item.','Observar se revisores localizam e corrigem a divergência; medir tempo e concordância.')])
sub('O que eu gostaria de ver melhor nos modelos Google')
p('Para este caso, a prioridade é aderência semântica ao sistema em análise: manter separados o componente instalado e o candidato; usar consistentemente os tipos de relação; distinguir valor nominal, faixa e limite; e vincular cada conclusão ao trecho que a sustenta. A segunda prioridade é estabilidade entre repetições. A terceira é reduzir a latência mantendo evidência verificável.')
p('As melhorias do lado do produto são igualmente necessárias. Um prompt mais claro, uma referência validada e um avaliador que reconheça fontes equivalentes podem evitar falsos alarmes e revelar melhor as falhas reais. A hipótese a testar é um fluxo híbrido: modelo para propor e localizar, verificações determinísticas para contratos e números, revisão humana para aceitação.')
sub('Próximo ciclo recomendado')
p('Congelar uma versão, ampliar documentos e cenários, repetir todas as tarefas por modelo e pré-definir critérios de sucesso. Medir qualidade estrutural, evidência e tempo separadamente. Nenhuma dessas novas execuções foi realizada para produzir este relatório.')

page();h('11. Anexo de métricas e definições')
p('As tabelas usam a média aritmética dos indicadores salvos por resposta, como a tela de comparação. Não usam F1 calculado a partir da soma de todas as arestas. “n” é o número de respostas efetivamente avaliadas naquela combinação. Erros técnicos, interrupções e Astra d4408d não entram no denominador de qualidade.')
short={'entity_extraction':'Extração','relationship_extraction':'Relações','change_impact':'Impactos','impact_explanation':'Explicações','one_hop':'Inferência'}
rows=[]
order={'gemini-3.1-pro-preview':0,'gpt-6-astra':1,'claude-fable-5-1':2}
for g in sorted(A,key=lambda x:(order[x['model']],x['comparison_hash'],list(short).index(x['task']),x['difficulty'])):
    name={'gemini-3.1-pro-preview':'Gemini','gpt-6-astra':'Astra','claude-fable-5-1':'Claude'}[g['model']]
    primary='entity_f1' if g['task']=='entity_extraction' else 'relationship_f1' if g['task']=='relationship_extraction' else 'impact_f1'
    rows.append([name+' '+g['comparison_hash'][:6],short[g['task']]+' '+g['difficulty'][:2],g['n'],pct(metric(g,primary)),pct(g['metrics'].get('explanation_rule_pass_rate',{}).get('mean')),num(g['latency_seconds']['mean'])])
table(['Modelo / versão','Tarefa','n','F1 principal','Justif. aceitas','Média s'],rows)
p('“F1 principal” muda conforme a tarefa: entidades, relações ou conjunto de requisitos afetados. A coluna de justificativas não é aplicável à extração nem às relações. Comparar colunas como se fossem uma nota geral de inteligência seria incorreto.')
sub('Definições de leitura')
p('Precisão = itens previstos que coincidem com a referência ÷ itens previstos. Cobertura (recall) = itens esperados encontrados ÷ itens esperados. F1 combina precisão e cobertura por média harmônica. Para relações, a correspondência usa origem, tipo e alvo. A versão “com suporte” exige também a evidência aceita.')
p('Consistência = proporção de pares de respostas com a mesma assinatura estruturada; não mede qualidade da prosa nem probabilidade de acerto futuro. “Não suportado” significa reprovado por uma regra do benchmark e pode incluir ausência na referência, citação em página não aceita, valor ou qualificador divergente.')

page();h('12. Proveniência e conferência')
p('Estudo e continuações: a66272eaf0c24305b11bd3c368db4798. Corte: 27/09/2026 às 22:33:24 BRT (28/09/2026 01:33:24 UTC). A base foi copiada para análise privada; o banco original e os resultados históricos não foram editados. Não houve commit, publicação externa nem chamada aos provedores durante a preparação.')
table(['Recorte auditado','Contagem'],[
('Registros de tarefas no histórico, incluindo fila/interrupções','123'),
('Respostas avaliadas no histórico','61'),
('Astra d4408d excluído','18'),
('Respostas elegíveis neste relatório','43 = 18 Gemini + 17 Astra + 8 Claude'),
('Erros técnicos explícitos','4, todos Claude: 3 contrato + 1 billing'),
('Registros interrompidos','58; não equivalem a 58 falhas de API')])
sub('Arquivos que acompanham o documento')
bullet('metricas_por_resposta.csv: 43 linhas de resultados elegíveis, com IDs, modelo, protocolo, tarefa, nível, métricas e latência.')
bullet('metricas_agregadas.json: 19 grupos de métricas com médias, mínimo/máximo, amostra e proveniência. Não contém chaves de API.')
bullet('figuras/: gráficos em PNG e SVG; relatorio.md: texto editável; resposta_formulario.txt: resposta sugerida à questão 5.')
sub('Rastreabilidade dos exemplos')
table(['Exemplo','ID da resposta ou chamada'],[
('Gemini · extração · repetição 1','b71e084e76be42dd88837e257b19e02c'),
('Gemini · grafo · repetição 1','397e93c4fdae4e6cbc82f617816b4c74'),
('Gemini · explicações L1 · repetição 1','03b1d0c62d8b47e5af68ff23e719f81b'),
('Claude · rejeição CHG-004/REQ-002','0f7a4ebe56c944d1a83bdabf308d02af'),
('Claude · billing','13d7ad57258c4b448d0526250aaf23bd')])
sub('Fontes primárias locais')
p('Métricas e respostas: data/benchmark.sqlite3, tabelas executions, live_calls e task_results. Regras: benchmark/evaluate.py e snapshots com evaluation_code_hash. Validade: config/result_exclusions.json. Limitações da referência: docs/ground_truth_review.md. Documentos de origem: registro de fontes e hashes dos PDFs preservados nos snapshots. Os números do relatório vêm das avaliações salvas; as explicações adicionais identificam limites sem alterar as notas.')
p('Fontes técnicas preservadas: especificações Noctua NF-A4x10 5V e NF-A4x20 5V; datasheets TI TMP117 (SNOSD82D), TPS22919 (SLVSEN5B) e TPS22917 (SLVSDW8B).')

doc.save(OUT/'NORTE_relatorio_Gama_Fund.docx')
(OUT/'relatorio.md').write_text('\n'.join(md),encoding='utf-8')
print('Saved DOCX and Markdown, 13 planned pages; 5 figures in PNG/SVG.')
